#!/usr/bin/env python3
"""Export a markdown story file to a formatted .docx file."""

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
except ImportError:
    print("Error: python-docx is required. Install it with: pip install python-docx", file=sys.stderr)
    sys.exit(1)


def set_font(run, name_cn="等线", name_en="Calibri", size=12, bold=False, color=None):
    """Set font properties for a run, supporting Chinese fonts."""
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run.font.name = name_en
    r = run._element
    rPr = r.find(qn('w:rPr'))
    if rPr is None:
        rPr = r.makeelement(qn('w:rPr'), {})
        r.insert(0, rPr)
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), name_cn)


def md_to_docx(md_path: str, output_path: str, title: str | None = None):
    """Convert a markdown file to a .docx document."""
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '等线')

    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    lines = content.split("\n")

    i = 0
    while i < len(lines):
        line = lines[i]

        # Title — h1
        if line.startswith("# ") and not title:
            title = line[2:].strip()
            heading = doc.add_heading(title, level=0)
            for run in heading.runs:
                set_font(run, size=22, bold=True)
            i += 1
            continue

        # Add document title at the top
        if i == 0 and title and line.startswith("# "):
            doc_title = doc.add_heading(title, level=0)
            for run in doc_title.runs:
                set_font(run, size=22, bold=True)
            i += 1
            continue
        elif i == 0 and title:
            doc_title = doc.add_heading(title, level=0)
            for run in doc_title.runs:
                set_font(run, size=22, bold=True)
            # don't skip this line, process it normally
        elif i == 0 and not title:
            if line.startswith("# "):
                title = line[2:].strip()
                heading = doc.add_heading(title, level=0)
                for run in heading.runs:
                    set_font(run, size=22, bold=True)
                i += 1
                continue

        # Headings
        if line.startswith("## "):
            h = doc.add_heading(line[3:].strip(), level=1)
            for run in h.runs:
                set_font(run, size=16, bold=True)
        elif line.startswith("### "):
            h = doc.add_heading(line[4:].strip(), level=2)
            for run in h.runs:
                set_font(run, size=14, bold=True)
        elif line.startswith("#### "):
            h = doc.add_heading(line[5:].strip(), level=3)
            for run in h.runs:
                set_font(run, size=13, bold=True)

        # Horizontal rule
        elif re.match(r"^[-*_]{3,}$", line.strip()):
            doc.add_paragraph("—" * 40)

        # Blockquote
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.style = doc.styles['Quote']
            run = p.add_run(line[2:].strip())
            set_font(run, size=11)

        # Unordered list
        elif re.match(r"^[-*+]\s", line):
            p = doc.add_paragraph(style='List Bullet')
            # Process inline formatting
            text = re.sub(r"^[-*+]\s+", "", line)
            _add_formatted_text(p, text)

        # Ordered list
        elif re.match(r"^\d+[.)]\s", line):
            text = re.sub(r"^\d+[.)]\s+", "", line)
            p = doc.add_paragraph(text, style='List Number')

        # Empty line
        elif not line.strip():
            doc.add_paragraph()

        # Normal paragraph
        else:
            p = doc.add_paragraph()
            _add_formatted_text(p, line)

        i += 1

    # Save
    doc.save(output_path)
    print(f"Exported to {output_path}")


def _add_formatted_text(paragraph, text: str):
    """Add text with inline markdown formatting (bold, italic) to a paragraph."""
    # Bold: **text**
    # Italic: *text*
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=12, bold=True)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, size=12)
            run.font.italic = True
        else:
            run = paragraph.add_run(part)
            set_font(run, size=12)


def main():
    parser = argparse.ArgumentParser(description="Export markdown story to .docx")
    parser.add_argument("--source", required=True, help="Path to source .md file")
    parser.add_argument("--output", required=True, help="Path to output .docx file")
    parser.add_argument("--title", default=None, help="Document title (defaults to first H1)")
    args = parser.parse_args()

    if not Path(args.source).exists():
        print(f"Error: source file not found: {args.source}", file=sys.stderr)
        sys.exit(1)

    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    md_to_docx(args.source, args.output, args.title)


if __name__ == "__main__":
    main()
